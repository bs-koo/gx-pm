"""Export real workbooks to check readable sizing without changing document data."""
import contextlib
import io
import tempfile
import unittest
from pathlib import Path

import openpyxl
from helpers import load_export_module


class ReadableWorkbookTest(unittest.TestCase):
    def setUp(self):
        self.mod = load_export_module()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def export(self, doc_type, body, set_index=0, draft=False):
        headers = self.mod.DOCUMENT_PROFILES[doc_type]['columns'][set_index]
        lines = ['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |']
        lines += ['| ' + ' | '.join(row) + ' |' for row in body]
        path = Path(self.tmp.name) / f'DEMO-{doc_type}.xlsx'
        out, _ = self.mod.create_xlsx([(f'DEMO-{doc_type}.md', [(doc_type, lines)])], str(path), draft=draft)
        wb = openpyxl.load_workbook(out)
        self.addCleanup(wb.close)
        return wb.active

    def test_confirmation_explanation_and_edit_fields_are_readable(self):
        reason = ('기존 ddl.sql의 TB_AUTH.AUTHRT_NM이 VARCHAR(100)이므로 저장 제약에 맞춰 최대 100자로 작성했습니다. '
                  'RFP에는 길이가 명시되어 있지 않아 이 제약을 유지할지 확인합니다.')
        ws = self.export('확인요청서', [['CQ-001', 'AN-03', '1', 'FN-003 권한명 길이',
                                      '100자로 제한', reason, '', '', '']], 1)
        self.assertEqual(ws['F2'].value, reason)
        self.assertGreater(ws.column_dimensions['F'].width, 60)
        self.assertGreater(ws.row_dimensions[2].height, 30)
        for cell in ('E2', 'F2', 'H2'):
            self.assertEqual(ws[cell].alignment.horizontal, 'left')
            self.assertTrue(ws[cell].alignment.wrap_text)
        self.assertEqual(ws['G2'].alignment.horizontal, 'center')
        self.assertIsNone(ws['G2'].value)
        self.assertGreaterEqual(ws.column_dimensions['H'].width, 24)
        self.assertEqual(ws['H2'].fill.fgColor.rgb[-6:], 'FFF9E3')

    def test_html_breaks_become_visible_lines_with_enough_height(self):
        ws = self.export('확인요청서', [['현재', '첫 번째 줄<br>두 번째 줄<BR/>세 번째 줄<br />마지막 줄']])
        self.assertEqual(ws['B2'].value, '첫 번째 줄\n두 번째 줄\n세 번째 줄\n마지막 줄')
        self.assertGreaterEqual(ws.row_dimensions[2].height, 66)
        self.assertLessEqual(ws.column_dimensions['B'].width, 104)

    def test_explicit_lines_do_not_add_up_to_column_width(self):
        self.assertEqual(self.mod._column_width(['가' * 10 + '\n' + '나' * 10]), 24)
        self.assertEqual(self.mod._display_width('e\u0301'), 1)

    def test_merged_narrative_has_enough_total_height(self):
        condition = '사전조건을 만족하는 계정과 데이터가 준비되어 있어야 합니다. ' * 18
        rows = [[f'UT-{n:03}', 'FN-001', 'REQ-001', '권한 등록', condition,
                 '입력값', '저장됨', '데이터 존재', '', '', '', ''] for n in range(1, 4)]
        ws = self.export('단위테스트계획서', rows)
        self.assertIn('E2:E4', [str(m) for m in ws.merged_cells.ranges])
        self.assertEqual(ws['E2'].value, condition.strip())
        needed = self.mod._text_height(condition.strip(), ws.column_dimensions['E'].width)
        self.assertGreaterEqual(sum(ws.row_dimensions[r].height for r in range(2, 5)) + 0.01, needed)
        self.assertEqual(ws['E2'].alignment.horizontal, 'left')
        self.assertTrue(all(ws.row_dimensions[r].height <= 409 for r in range(2, 5)))

    def test_extreme_text_is_preserved_and_size_limit_reported(self):
        original = '아주 긴 근거 문장을 삭제하지 않습니다. ' * 600
        log = io.StringIO()
        with contextlib.redirect_stderr(log):
            ws = self.export('확인요청서', [['현재', original]])
        self.assertEqual(ws['B2'].value, original.strip())
        self.assertEqual(ws.row_dimensions[2].height, 409)
        self.assertIn('B2', log.getvalue())
        self.assertIn('상한', log.getvalue())

    def test_print_headers_and_draft_banner_do_not_hide_data(self):
        for draft in (False, True):
            with self.subTest(draft=draft):
                ws = self.export('확인요청서', [['현재', '설명입니다.']], draft=draft)
                offset = int(draft)
                self.assertEqual(ws[f'A{2 + offset}'].value, '현재')
                self.assertEqual(ws.freeze_panes, f'A{2 + offset}')
                self.assertEqual(ws.print_title_rows, f'$1:${1 + offset}')
                self.assertGreaterEqual(ws.row_dimensions[1].height, 30)
                self.assertEqual(ws[f'B{2 + offset}'].font.sz, 10)


if __name__ == '__main__':
    unittest.main()
