"""Real XLSX round trips and prompt contracts for document review workflows."""
import contextlib
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from helpers import PLUGIN_ROOT, load_export_module


class ConfirmationWorkflowTest(unittest.TestCase):
    def test_five_origins_export_with_responses_and_history_intact(self):
        mod = load_export_module()
        profile = mod.DOCUMENT_PROFILES['확인요청서']
        for origin, command in [('AN-02','요구사항정의서'), ('AN-03','기능명세서'), ('DE-08','테이블정의서'),
                                ('DE-13','단위테스트계획서'), ('AN-05','추적매트릭스')]:
            with self.subTest(origin=origin), tempfile.TemporaryDirectory() as tmp:
                rows = [
                    [['현재', '1차 /gx-'+command]],
                    [['CQ-001',origin,'1','REQ-001 범위','제안 범위','원문 충돌','','','']],
                    [['CQ-002',origin,'1','REQ-002 기준값','기준값','근거 없음','확정','30','2026-10-26']],
                    [['1','2026-10-26','/gx-'+command,'CQ-001',origin,'REQ-001 범위','범위 확인','미응답','이월']],
                ]
                tables = []
                for title, headers, body in zip(profile['sheet_names'], profile['columns'], rows):
                    lines = ['| '+' | '.join(headers)+' |', '| '+' | '.join(['---']*len(headers))+' |']
                    lines += ['| '+' | '.join(row)+' |' for row in body]
                    tables.append((title,lines))
                target = str(Path(tmp)/'DEMO-확인요청서.xlsx')
                mod.create_xlsx([('DEMO-확인요청서.md',tables)],target)
                wb = openpyxl.load_workbook(target)
                self.addCleanup(wb.close)
                self.assertEqual(wb.sheetnames, profile['sheet_names'])
                self.assertEqual(wb['가정 확인']['B2'].value,origin)
                self.assertIsNone(wb['가정 확인']['G2'].value, 'No preselected acceptance')
                self.assertEqual(wb['가정 확인']['G2'].fill.fgColor.rgb[-6:],'FFF9E3')
                self.assertEqual(wb['미확정 확인']['H2'].value,'30')
                self.assertEqual(wb['요청 이력']['H2'].value,'미응답')
                self.assertEqual(wb['요청 이력']['I2'].value,'이월')

    def test_shared_skill_protects_answers_and_waits_for_real_input(self):
        text=(PLUGIN_ROOT/'skills/review-confirmations/SKILL.md').read_text(encoding='utf-8')
        for required in ('AN-02','AN-03','DE-08','DE-13','AN-05','실제 응답을 기다린다','충돌','백업','파일이 없으면','차수·물은 횟수'):
            self.assertIn(required,text)
        common=(PLUGIN_ROOT/'templates/approval-protocol.md').read_text(encoding='utf-8')
        for required in ('진행 안내·파일 경로·다음 제안은 질문이 아니다','이미 정한','비동기','일부 답만','자유 입력'):
            self.assertIn(required,common)

    def test_alias_delegates_without_copying_business_rules(self):
        alias=(PLUGIN_ROOT/'commands/gx-단위테스트결과반영.md').read_text(encoding='utf-8')
        self.assertIn('commands/gx-단위테스트실행.md',alias)
        self.assertNotIn('### Step',alias)
        self.assertIn('별도 승인을 추가하지 않는다',alias)


class DraftExportTest(unittest.TestCase):
    def setUp(self):
        self.mod=load_export_module()
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.text='''# 검토

## 개정이력
| 버전 | 개정일 | 개정 사유 | 개정 내용 | 작성자 | 승인자 |
|---|---|---|---|---|---|
| 1.0 | 2026.10.26 | 신규 | 검토용 | 담당자 | |

## 내용
| 요구사항ID | 요구사항명 | 상세내용 | 유형 |
|---|---|---|---|
| REQ-001 | 기준 확인 | [미확정] | 기능 |
'''

    def test_draft_preserves_final_and_marks_every_sheet(self):
        final=self.root/'DEMO-요구사항정의서.xlsx'
        final.write_bytes(b'existing final must survive')
        tables=self.mod.parse_markdown_tables(self.text)
        out,count=self.mod.create_xlsx([('DEMO-요구사항정의서.md',tables)],str(final),draft=True)
        self.assertEqual(Path(out).name,'DEMO-요구사항정의서-검토용초안.xlsx')
        self.assertEqual(final.read_bytes(),b'existing final must survive')
        wb=openpyxl.load_workbook(out)
        self.addCleanup(wb.close)
        self.assertEqual(count,2)
        self.assertIn('검토용초안',wb.properties.title)
        for ws in wb:
            self.assertIn('검토용초안',ws['A1'].value)
            self.assertIn('검토용초안',ws.oddHeader.center.text)
            self.assertEqual(ws.freeze_panes,'A3')
            self.assertTrue(ws.auto_filter.ref.startswith('A2:'))

    def test_draft_cli_separate_and_combined_use_safe_names(self):
        md=self.root/'DEMO-요구사항정의서.md'
        md.write_text(self.text,encoding='utf-8')
        for separate in (True,False):
            output=self.root/('xlsx' if separate else 'combined.xlsx')
            args=['export-xlsx.py','--draft','--output',str(output),str(md)]
            if separate:
                args.append('--separate')
            with self.subTest(separate=separate),patch.object(sys,'argv',args),contextlib.redirect_stdout(io.StringIO()):
                self.mod.main()
            expected=output/'DEMO-요구사항정의서-검토용초안.xlsx' if separate else self.root/'combined-검토용초안.xlsx'
            self.assertTrue(expected.is_file())

    def test_already_draft_name_is_not_doubled(self):
        path=self.root/'검토-검토용초안.xlsx'
        out,_=self.mod.create_xlsx([('demo',self.mod.parse_markdown_tables(self.text))],str(path),draft=True)
        self.assertEqual(Path(out),path)

    def test_draft_banner_shifts_validation_and_input_cells(self):
        profile=self.mod.DOCUMENT_PROFILES['확인요청서']
        headers=profile['columns'][1]
        lines=['| '+' | '.join(headers)+' |', '| '+' | '.join(['---']*len(headers))+' |',
               '| CQ-001 | DE-08 | 1 | TB_A.NAME | 100 | 기준 부재 | | | |']
        out,_=self.mod.create_xlsx([('DEMO-확인요청서.md',[('가정 확인',lines)])],
                                   str(self.root/'DEMO-확인요청서.xlsx'),draft=True)
        wb=openpyxl.load_workbook(out)
        self.addCleanup(wb.close)
        ws=wb['가정 확인']
        self.assertEqual(ws['A3'].value,'CQ-001')
        self.assertEqual(ws['G2'].value,'판정')
        self.assertIsNone(ws['G3'].value)
        self.assertEqual(ws['G3'].fill.fgColor.rgb[-6:],'FFF9E3')
        self.assertEqual([str(dv.sqref) for dv in ws.data_validations.dataValidation],['G3'])

    def test_stdin_draft_marks_empty_fallback(self):
        target=self.root/'stdin.xlsx'
        with patch.object(sys,'argv',['export-xlsx.py','--draft','--output',str(target)]), \
                patch.object(sys,'stdin',io.StringIO('')),contextlib.redirect_stdout(io.StringIO()), \
                contextlib.redirect_stderr(io.StringIO()):
            self.mod.main()
        self.assertFalse(target.exists())
        wb=openpyxl.load_workbook(self.root/'stdin-검토용초안.xlsx')
        self.addCleanup(wb.close)
        self.assertIn('검토용초안',wb.active['A1'].value)
        self.assertIn('검토용초안',wb.active.oddHeader.center.text)


if __name__=='__main__':
    unittest.main()
